"""Actual bounded native T01 integration proof; failures remain visible in JSON."""

import hashlib
import importlib.metadata
import importlib.util
import io
import json
import math
import os
from pathlib import Path
import platform
import resource
import shutil
import subprocess
import tempfile
import time

from docx import Document
from defusedxml.ElementTree import fromstring
from defusedxml.common import EntitiesForbidden
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont, features

from sandbox_proof import run as sandbox_run


def command(args, timeout=30, include_stderr=False):
    result = subprocess.run(args, capture_output=True, text=True, timeout=timeout)
    if result.returncode:
        raise RuntimeError(f"{args[0]} exit={result.returncode}: {result.stderr[:2000]}")
    return result.stdout + (result.stderr if include_stderr else "")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf(path, scan, mixed):
    """Small real PDF: same-page image XObject and optional selectable text."""
    image_bytes = io.BytesIO()
    scan.save(image_bytes, format="JPEG", quality=95)
    data = image_bytes.getvalue()
    content = b"q 500 0 0 180 45 260 cm /Im0 Do Q\n"
    if mixed:
        content += b"BT /F1 22 Tf 45 690 Td (SELECTABLE REGION ORANGE 7319) Tj ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> /XObject << /Im0 5 0 R >> >> /Contents 6 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Type /XObject /Subtype /Image /Width {scan.width} /Height {scan.height} /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"endstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode())
    path.write_bytes(result)


def main(fixture_root=Path("/fixtures"), inside_sandbox=False):
    started = time.monotonic()
    report = {"scope": "T01 synthetic native proof only; not release acceptance", "inside_sandbox": inside_sandbox, "architecture": platform.machine(), "python": platform.python_version(), "cases": {}}
    report["packages"] = {name: importlib.metadata.version(name) for name in ("Pillow", "python-docx", "lxml", "openpyxl", "defusedxml")}
    if not inside_sandbox:
        report["dpkg"] = Path("/opt/dpkg-proof.tsv").read_text().splitlines()
        report["pip"] = Path("/opt/pip-proof.txt").read_text().splitlines()
    report["hevc_fixture"] = json.loads((fixture_root / "manifest.json").read_text())
    with tempfile.TemporaryDirectory(prefix="t01-", dir="/tmp") as temporary:
        work = Path(temporary)
        originals = {}

        def check(name, fn):
            before = time.monotonic()
            try:
                report["cases"][name] = {"status": "pass", "evidence": fn(), "wall_seconds": round(time.monotonic() - before, 3)}
            except Exception as error:
                report["cases"][name] = {"status": "fail", "error": str(error), "wall_seconds": round(time.monotonic() - before, 3)}

        def native_versions():
            decoders = command(["heif-dec", "--list-decoders"])
            loader = command(["ldd", "/usr/local/lib/libheif.so.1"])
            assert "libde265" in decoders and "libde265.so.0" in loader
            assert not any(name in loader for name in ("libx265", "libaom", "libdav1d", "libavcodec", "libx264"))
            assert not any(name in decoders for name in ("dav1d", "aom =", "ffmpeg"))
            assert importlib.util.find_spec("pillow_heif") is None, "encoder must not ship in final image"
            security = json.loads(command(["codec-security-proof", str(fixture_root / "synthetic-hevc.heic")]))
            assert security["global_limits_finite"] and security["tightened_100_pixel_limit_rejected"]
            return {"heif_info": command(["heif-info", "--help"], include_stderr=True), "heif_dec": command(["heif-dec", "--help"], include_stderr=True), "decoders": decoders, "heif_loader": command(["ldd", "/usr/local/bin/heif-dec"]), "libheif_loader": loader, "security_limits": security, "tesseract": command(["tesseract", "--version"]), "languages": command(["tesseract", "--list-langs"]), "bubblewrap": command(["bwrap", "--version"])}

        check("native_versions_and_linkage", native_versions)

        def raster():
            results = {}
            for fmt, feature in (("JPEG", "jpg"), ("PNG", "zlib"), ("WEBP", "webp")):
                assert features.check(feature), feature
                stream = io.BytesIO()
                Image.new("RGB", (80, 60), "red").save(stream, format=fmt)
                raw = stream.getvalue()
                with Image.open(io.BytesIO(raw)) as image:
                    image.load()
                    assert image.size == (80, 60) and image.getpixel((30, 30))[0] > 220
                results[fmt] = {"bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
            return results

        check("in_memory_jpeg_png_webp", raster)

        def hevc():
            results = {}
            for suffix in ("heic", "heif"):
                source = fixture_root / f"synthetic-hevc.{suffix}"
                originals[str(source)] = sha(source)
                info = command(["heif-info", str(source)])
                derivatives = {}
                for extension in ("png", "jpg"):
                    output = work / f"converted-{suffix}.{extension}"
                    conversion = command(["heif-dec", "-d", "libde265", str(source), str(output)])
                    with Image.open(output) as image:
                        image.load()
                        assert image.size == (640, 480)
                        red, blue = image.convert("RGB").getpixel((100, 200)), image.convert("RGB").getpixel((450, 200))
                        assert red[0] > 200 and red[2] < 50 and blue[2] > 200 and blue[0] < 50
                    derivatives[extension] = {"conversion": conversion, "bytes": output.stat().st_size}
                assert sha(source) == originals[str(source)]
                assert source.read_bytes()[4:8] == b"ftyp"
                assert b"hvc1" in source.read_bytes(), "fixture must contain actual HEVC item"
                results[suffix] = {"info": info, "derivatives": derivatives, "decoder": "explicit -d libde265", "original_sha256": sha(source), "original_unchanged": True}
            return results

        check("libheif_hevc_decode", hevc)

        scan = Image.new("RGB", (1400, 500), "white")
        draw = ImageDraw.Draw(scan)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 56)
        draw.text((30, 80), "SCANNED REGION VIOLET 4826", font=font, fill="black")
        draw.text((30, 210), "DOCUMENTO ESPAÑOL NÚMERO 5937", font=font, fill="black")

        def ocr(mixed):
            source = work / ("mixed.pdf" if mixed else "scan.pdf")
            pdf(source, scan, mixed)
            originals[str(source)] = sha(source)
            selectable = command(["pdftotext", str(source), "-"])
            bbox = command(["pdftotext", "-bbox-layout", str(source), "-"])
            if mixed:
                assert "SELECTABLE REGION ORANGE 7319" in selectable
                assert "SCANNED REGION" not in selectable
            else:
                assert not selectable.strip()
            prefix = work / ("mixed-page" if mixed else "scan-page")
            command(["pdftoppm", "-f", "1", "-singlefile", "-r", "150", "-png", str(source), str(prefix)])
            rendered = prefix.with_suffix(".png")
            # Poppler discovers image boxes, including beside nonempty text.
            # This proves simple raster-XObject region handling, not arbitrary
            # vector scans, overlapping layers or production OCR deduplication.
            layout = fromstring(command(["pdftohtml", "-xml", "-hidden", "-zoom", "1", "-stdout", str(source)]))
            boxes = []
            ocr_texts = []
            with Image.open(rendered) as page:
                assert page.width * page.height <= 8_000_000
                for index, element in enumerate(layout.findall("./page/image")):
                    left, top, width, height = (float(element.attrib[key]) for key in ("left", "top", "width", "height"))
                    box = (max(0, math.floor(left * 150 / 72) - 8), max(0, math.floor(top * 150 / 72) - 8), min(page.width, math.ceil((left + width) * 150 / 72) + 8), min(page.height, math.ceil((top + height) * 150 / 72) + 8))
                    crop = work / f"{source.stem}-region-{index}.png"
                    page.crop(box).save(crop)
                    boxes.append({"page": 1, "pdf_image_box": [left, top, width, height], "raster_crop": box})
                    ocr_texts.append(command(["tesseract", str(crop), "stdout", "-l", "eng+spa", "--psm", "6"], timeout=15))
            assert boxes, "no scanned image regions discovered"
            text = "\n".join(ocr_texts)
            assert "VIOLET 4826" in text and "5937" in text, text
            assert sha(source) == originals[str(source)]
            return {"selectable": selectable.strip(), "ocr_region": text.strip(), "selectable_bbox_present": "<word" in bbox, "regions": boxes, "merged_text": (selectable + "\n" + text).strip(), "original_sha256": sha(source), "original_unchanged": True, "automatic_region_detection": "Poppler image XObject boxes", "overlap_deduplication": "not exercised"}

        check("scanned_pdf_ocr", lambda: ocr(False))
        check("same_page_selectable_and_scan_ocr", lambda: ocr(True))

        def documents():
            docx = work / "synthetic.docx"
            doc = Document()
            doc.add_paragraph("DOCX PARAGRAPH 2648")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "DOCX TABLE 3759"
            doc.sections[0].header.paragraphs[0].text = "HEADER 4860"
            doc.save(docx)
            originals[str(docx)] = sha(docx)
            parsed = Document(docx)
            assert parsed.paragraphs[0].text == "DOCX PARAGRAPH 2648"
            assert parsed.tables[0].cell(0, 0).text == "DOCX TABLE 3759"
            assert parsed.sections[0].header.paragraphs[0].text == "HEADER 4860"
            xlsx = work / "synthetic.xlsx"
            book = Workbook()
            book.active["A1"] = "XLSX VALUE 5971"
            book.active["B1"] = "=1+2"
            book.create_sheet("Hidden").sheet_state = "hidden"
            book["Hidden"]["A1"] = "HIDDEN VALUE 6082"
            book.save(xlsx)
            book.close()
            originals[str(xlsx)] = sha(xlsx)
            parsed_book = load_workbook(xlsx, read_only=True, data_only=False, keep_links=False)
            assert parsed_book.active["A1"].value == "XLSX VALUE 5971"
            assert parsed_book.active["B1"].value == "=1+2"
            assert parsed_book["Hidden"]["A1"].value == "HIDDEN VALUE 6082"
            assert parsed_book["Hidden"].sheet_state == "hidden"
            parsed_book.close()
            cached = load_workbook(xlsx, read_only=True, data_only=True, keep_links=False)
            assert cached.active["B1"].value is None
            cached.close()
            for path in (docx, xlsx):
                assert sha(path) == originals[str(path)]
            return {"docx": "paragraph/table/header", "xlsx": "value/hidden sheet/literal formula/missing cache", "originals_unchanged": True}

        check("docx_xlsx", documents)

        def xml():
            assert fromstring("<root><text>XML VALUE 7193</text></root>").find("text").text == "XML VALUE 7193"
            try:
                fromstring('<!DOCTYPE x [<!ENTITY unsafe SYSTEM "file:///private-coordinator/sentinel">]><x>&unsafe;</x>')
            except EntitiesForbidden:
                return {"text_parse": True, "external_entity": "rejected before resolution"}
            raise AssertionError("external XML entity accepted")

        check("xml_external_entity_rejection", xml)

        def sandbox():
            job = work / "job"
            job.mkdir()
            (job / "tmp").mkdir()
            shutil.copytree(fixture_root, job / "fixtures")
            evidence = sandbox_run(job, scratch_root=work)
            if evidence["exit"]:
                raise RuntimeError(json.dumps(evidence))
            assert (job / "child-result.txt").is_file()
            return evidence

        if not inside_sandbox:
            from test_sandbox_paths import run_path_tests

            check("sandbox_path_source_validation", run_path_tests)
            check("rootless_no_proc_seccomp_sandbox", sandbox)
        report["scratch_bytes_end"] = sum(path.stat().st_size for path in work.rglob("*") if path.is_file())
        report["original_hashes_unchanged"] = all(sha(Path(path)) == digest for path, digest in originals.items())
    report["wall_seconds"] = round(time.monotonic() - started, 3)
    report["max_rss_kib"] = {"python_self": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss, "max_child_not_sum": resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss}
    report["cpu_seconds"] = {"self": resource.getrusage(resource.RUSAGE_SELF).ru_utime + resource.getrusage(resource.RUSAGE_SELF).ru_stime, "children": resource.getrusage(resource.RUSAGE_CHILDREN).ru_utime + resource.getrusage(resource.RUSAGE_CHILDREN).ru_stime}
    print(json.dumps(report, indent=2))
    return 0 if all(case["status"] == "pass" for case in report["cases"].values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
